"""Parser e Document IR: coordinate, ruoli di sezione, tabelle (PRD 9.1-9.2)."""

from __future__ import annotations

from pathlib import Path

import pytest
from conftest import ProjectFactory

from ntruth.ingest.project import Project
from ntruth.parsers.registry import build_document_ir
from ntruth.parsers.sections import classify_heading
from ntruth.schemas.document import ParserStatus, SectionRole
from ntruth.schemas.manifest import ReleaseProfile


@pytest.mark.parametrize(
    ("heading", "role"),
    [
        ("Materials and Methods", SectionRole.METHODS),
        ("Statistical analysis", SectionRole.STATISTICS),
        ("Analisi statistica", SectionRole.STATISTICS),
        ("Figure 3", SectionRole.FIGURE_LEGEND),
        ("Supplementary Table 2", SectionRole.TABLE_CAPTION),
        ("Cell culture", SectionRole.SAMPLE_DESCRIPTION),
        ("Exclusion criteria", SectionRole.EXCLUSION),
        ("References", SectionRole.REFERENCES),
        ("Una sezione senza nome noto", SectionRole.OTHER),
    ],
)
def test_heading_classification(heading: str, role: SectionRole) -> None:
    assert classify_heading(heading)[0] is role


def test_paragraph_offsets_point_at_the_original_text(make_project: ProjectFactory) -> None:
    """Il Document IR conserva coordinate verificabili (PRD 11.3)."""
    project = make_project({"methods.md": "# Methods\n\nCells were treated.\n\nn = 12 wells.\n"})
    ir = build_document_ir(project)
    assert ir.paragraphs
    for paragraph in ir.paragraphs:
        assert ir.texts[paragraph.file_id][paragraph.start : paragraph.end] == paragraph.text


def test_subsection_inherits_the_parent_role(make_project: ProjectFactory) -> None:
    """Una sottosezione con titolo di dominio resta dentro Methods."""
    project = make_project(
        {"m.md": "# Materials and Methods\n\n## Diet\n\nThe diet was assigned per cage.\n"}
    )
    ir = build_document_ir(project)
    diet = next(s for s in ir.sections if s.title == "Diet")
    assert diet.role is SectionRole.METHODS
    assert diet.role_origin == "inherited_from_parent_section"


def test_file_without_headings_is_declared_not_guessed(make_project: ProjectFactory) -> None:
    """FR-006: nessun fallimento silenzioso, la scelta di fallback e dichiarata."""
    project = make_project({"plain.txt": "Cells were treated with drug or vehicle.\n"})
    ir = build_document_ir(project)
    source = ir.files[0]
    assert any("nessuna intestazione" in w for w in source.warnings)
    assert all(s.role_confidence <= 0.4 for s in ir.sections)


def test_csv_is_parsed_with_columns_and_rows(make_project: ProjectFactory) -> None:
    project = make_project({"s.csv": "donor,well,treatment\nD1,W1,drug\nD1,W2,vehicle\n"})
    ir = build_document_ir(project)
    table = ir.tables[0]
    assert table.columns == ("donor", "well", "treatment")
    assert len(table.rows) == 2
    assert table.cell(0, "treatment") == "drug"


def test_unsupported_extension_is_rejected_with_a_reason(tmp_path: Path) -> None:
    source = tmp_path / "src"
    source.mkdir()
    (source / "note.rtf").write_text("x", encoding="utf-8")
    project = Project.create(tmp_path / "prj", name="t")
    result = project.add(source)
    assert not result.accepted
    assert result.rejected[0].reason is not None
    assert "estensione non supportata" in result.rejected[0].reason


def test_empty_project_produces_an_empty_but_valid_ir(tmp_path: Path) -> None:
    project = Project.create(tmp_path / "prj", name="vuoto")
    ir = build_document_ir(project)
    assert ir.files == ()
    assert ir.design_text() == []


def test_markdown_table_becomes_a_table(make_project: ProjectFactory) -> None:
    project = make_project(
        {"m.md": "# Methods\n\n| donor | well |\n| --- | --- |\n| D1 | W1 |\n| D2 | W2 |\n"}
    )
    ir = build_document_ir(project)
    assert ir.tables and ir.tables[0].columns == ("donor", "well")
    assert ir.files[0].status is ParserStatus.OK


def test_jats_sections_and_legends(make_project: ProjectFactory) -> None:
    xml = (
        "<article><front><article-meta><title-group><article-title>Study</article-title>"
        "</title-group></article-meta></front><body>"
        "<sec><title>Methods</title><p>Cells were treated with drug or vehicle.</p>"
        "<sec><title>Statistical analysis</title><p>n = 6 donors.</p></sec></sec>"
        "</body></article>"
    )
    project = make_project(
        {"a.xml": xml},
        release_profile=ReleaseProfile.EXTENDED_EXPERIMENTAL,
    )
    ir = build_document_ir(project)
    roles = {s.role for s in ir.sections}
    assert SectionRole.METHODS in roles
    assert SectionRole.STATISTICS in roles


def _write_xlsx(path: Path, *, cached: bool) -> None:
    """Scrive un XLSX con formula ``=1+1`` e, opzionalmente, la cache di Excel."""

    import shutil
    import zipfile

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "wells"
    sheet["A2"] = "=1+1"
    workbook.save(path)
    if not cached:
        return
    # Excel salva <f>1+1</f><v>2</v>: openpyxl scrive solo la formula, quindi
    # la cache viene iniettata nel XML del foglio per simulare il file reale.
    injectable = path.with_suffix(".cached.xlsx")
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(injectable, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for member in source.namelist():
            data = source.read(member)
            if member.endswith("sheet1.xml"):
                data = data.replace(b"<f>1+1</f>", b"<f>1+1</f><v>2</v>")
            target.writestr(member, data)
    shutil.move(injectable, path)


def test_xlsx_formula_uses_cached_value_when_available(tmp_path: Path) -> None:
    from ntruth.parsers.tabular import XlsxParser

    path = tmp_path / "sheet.xlsx"
    _write_xlsx(path, cached=True)
    doc = XlsxParser().parse(path)
    assert doc.tables, doc.warnings
    rows = [list(row.values()) for row in doc.tables[0].rows]
    assert rows == [["2"]]
    assert not any("formule senza valore" in warning for warning in doc.warnings)


def test_xlsx_formula_without_cache_stays_inert_and_warns(tmp_path: Path) -> None:
    from ntruth.parsers.tabular import XlsxParser

    path = tmp_path / "sheet.xlsx"
    _write_xlsx(path, cached=False)
    doc = XlsxParser().parse(path)
    assert doc.tables, doc.warnings
    rows = [list(row.values()) for row in doc.tables[0].rows]
    assert rows == [["'=1+1"]]
    assert any("formule senza valore" in warning for warning in doc.warnings)


def test_no_usable_output_message_includes_per_file_warnings() -> None:
    """L'abort riporta gli hint per-file del Document IR (audit 2026-09-05, F3)."""

    from ntruth.pipeline import _no_usable_output_message
    from ntruth.schemas.document import DocumentIR, SourceFile

    scanned = SourceFile(
        id="F-1",
        filename="scanned.pdf",
        relative_path="scanned.pdf",
        media_type="application/pdf",
        size_bytes=10,
        sha256="0" * 64,
        parser="pdf",
        parser_version="0.0.0",
        status=ParserStatus.FAILED,
        warnings=("nessun testo estraibile: PDF probabilmente scansionato, serve OCR esplicito",),
    )
    document = DocumentIR(id="IR-1", files=(scanned,))
    message = _no_usable_output_message(document)
    assert "scanned.pdf" in message
    assert "serve OCR esplicito" in message
